from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from nps_lens.services.classification_protocol import category_catalog


def taxonomy_excel(taxonomy: dict[str, Any]) -> bytes:
    """Export the canonical category criteria used by classification."""
    categories = category_catalog(taxonomy)
    if not categories:
        raise ValueError("No hay una taxonomía disponible para descargar.")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Taxonomía"
    sheet.append(["Palanca", "Subpalanca", "Criterio de clasificación"])
    for category in categories.values():
        sheet.append([category["lever"], category["sublever"], category["criterion"]])
    for row in sheet:
        for cell in row:
            # Taxonomy labels and criteria are text, including a leading '='.
            cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="10194F")
    for column, width in {"A": 32, "B": 38, "C": 100}.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
