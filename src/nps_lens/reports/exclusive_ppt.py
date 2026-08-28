from __future__ import annotations

import math
import sys
from contextlib import suppress
from dataclasses import dataclass
from datetime import date
from io import BytesIO
from pathlib import Path

import pandas as pd
import plotly.graph_objects as go
import plotly.io as pio
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt

from nps_lens.analytics.drivers import compute_nps_from_scores, grouped_driver_stats
from nps_lens.analytics.opportunities import Opportunity, rank_opportunities
from nps_lens.design.tokens import bbva_typography_tokens
from nps_lens.domain.normalization import EquivalenceRegistry
from nps_lens.reports.executive_ppt import BusinessPptResult


def find_exclusive_template_path() -> Path:
    """Resolve the bundled template in source and packaged desktop builds."""
    candidates = [
        Path.cwd() / "assets" / "ppt" / "templates" / "nuevo-informe-bbva.pptx",
        Path(__file__).resolve().parents[3]
        / "assets"
        / "ppt"
        / "templates"
        / "nuevo-informe-bbva.pptx",
    ]
    bundle_root = getattr(sys, "_MEIPASS", None)
    if bundle_root:
        candidates.insert(
            0,
            Path(str(bundle_root)) / "assets" / "ppt" / "templates" / "nuevo-informe-bbva.pptx",
        )
    for candidate in candidates:
        if candidate.exists():
            return candidate
    raise ValueError("No se encuentra la plantilla corporativa del informe exclusivo.")


def find_bei_logo_path() -> Path:
    candidates = [
        Path.cwd() / "frontend" / "public" / "assets" / "brand" / "bbva-bei.png",
        Path(__file__).resolve().parents[3]
        / "frontend"
        / "public"
        / "assets"
        / "brand"
        / "bbva-bei.png",
    ]
    bundle_root = getattr(sys, "_MEIPASS", None)
    if bundle_root:
        candidates.insert(
            0,
            Path(str(bundle_root)) / "assets" / "brand" / "bbva-bei.png",
        )
    for candidate in candidates:
        if candidate.exists():
            return candidate
    raise ValueError(
        "No se encuentra el logotipo oficial de BBVA Banca de Empresas e Instituciones."
    )


@dataclass(frozen=True)
class ExclusiveReportContext:
    service_origin: str
    service_origin_n1: str
    service_origin_n2: str
    period_start: date
    period_end: date
    helix_base_url: str


def _fmt(value: object, decimals: int = 1) -> str:
    try:
        number = float(value)
    except Exception:
        return "n/d"
    if not math.isfinite(number):
        return "n/d"
    return f"{number:,.{decimals}f}".replace(",", "_").replace(".", ",").replace("_", ".")


def _set_shape_text(shape: object, text: str) -> None:
    if not getattr(shape, "has_text_frame", False):
        return
    frame = shape.text_frame
    paragraphs = list(frame.paragraphs)
    template_run = next((run for paragraph in paragraphs for run in paragraph.runs), None)
    frame.clear()
    paragraph = frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = str(text)
    if template_run is not None:
        run.font.name = template_run.font.name
        run.font.size = template_run.font.size
        run.font.bold = template_run.font.bold
        run.font.italic = template_run.font.italic
        if template_run.font.color.type is not None:
            with suppress(AttributeError, TypeError):
                run.font.color.rgb = template_run.font.color.rgb
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE


def _set_mixed_statement(shape: object, headline: str, detail: str) -> None:
    frame = shape.text_frame
    source_runs = [run for paragraph in frame.paragraphs for run in paragraph.runs]
    frame.clear()
    paragraph = frame.paragraphs[0]
    for index, text in enumerate([headline, ": ", detail]):
        run = paragraph.add_run()
        run.text = text
        source = source_runs[min(index, len(source_runs) - 1)] if source_runs else None
        if source is not None:
            run.font.name = source.font.name
            run.font.size = source.font.size
            run.font.bold = source.font.bold
            if source.font.color.type is not None:
                with suppress(AttributeError, TypeError):
                    run.font.color.rgb = source.font.color.rgb
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE


def _set_card(shape: object, title: str, detail: str) -> None:
    frame = shape.text_frame
    frame.clear()
    title_run = frame.paragraphs[0].add_run()
    title_run.text = title
    title_run.font.name = "Benton Sans BBVA"
    title_run.font.size = Pt(14)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(7, 14, 70)
    detail_paragraph = frame.add_paragraph()
    detail_paragraph.space_before = Pt(96)
    detail_run = detail_paragraph.add_run()
    detail_run.text = detail
    detail_run.font.name = "Benton Sans BBVA"
    detail_run.font.size = Pt(11)
    detail_run.font.color.rgb = RGBColor(0, 19, 145)
    frame.word_wrap = True
    frame.auto_size = None


def _set_opportunity_metric(shape: object, *, uplift: float, confidence: float, n: int) -> None:
    frame = shape.text_frame
    frame.clear()
    metric = frame.paragraphs[0].add_run()
    metric.text = f"+{_fmt(uplift)} pts"
    metric.font.name = "Benton Sans BBVA"
    metric.font.size = Pt(25)
    metric.font.bold = True
    metric.font.color.rgb = RGBColor(0, 0, 0)
    evidence = frame.add_paragraph().add_run()
    evidence.text = f"Confianza {_fmt(confidence)} · n={n:,}".replace(",", ".")
    evidence.font.name = "Benton Sans BBVA"
    evidence.font.size = Pt(9)
    evidence.font.color.rgb = RGBColor(7, 14, 70)
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE


def _monthly_metrics(frame: pd.DataFrame) -> pd.DataFrame:
    if frame.empty:
        return pd.DataFrame(columns=["period", "nps", "detractor_rate"])
    work = frame.copy()
    work["Fecha"] = pd.to_datetime(work["Fecha"], errors="coerce")
    work["NPS"] = pd.to_numeric(work["NPS"], errors="coerce")
    work = work.dropna(subset=["Fecha", "NPS"])
    work["period"] = work["Fecha"].dt.to_period("M").dt.to_timestamp()
    rows = []
    for period, group in work.groupby("period", sort=True):
        rows.append(
            {
                "period": period,
                "nps": compute_nps_from_scores(group["NPS"]),
                "detractor_rate": float((group["NPS"] <= 6).mean()),
            }
        )
    return pd.DataFrame(rows)


def _chart_image(metrics: pd.DataFrame) -> bytes:
    figure = go.Figure()
    figure.add_trace(
        go.Scatter(
            x=metrics["period"],
            y=metrics["nps"],
            mode="lines+markers+text",
            text=[_fmt(value) for value in metrics["nps"]],
            textposition="top center",
            line={"color": "#001391", "width": 3},
            marker={"color": "#85c8ff", "size": 8},
            hoverinfo="skip",
        )
    )
    figure.update_layout(
        margin={"l": 45, "r": 18, "t": 24, "b": 42},
        paper_bgcolor="white",
        plot_bgcolor="white",
        showlegend=False,
        font={"family": "BentonSansBBVA Book", "color": "#070e46", "size": 12},
        xaxis={"showgrid": False, "tickformat": "%b %Y"},
        yaxis={"gridcolor": "#e2e6ea", "zerolinecolor": "#adb8c2", "title": "NPS clásico"},
    )
    return pio.to_image(figure, format="png", width=1100, height=620, scale=2)


def _replace_picture(slide: object, shape_index: int, image_bytes: bytes) -> None:
    shape = slide.shapes[shape_index]
    left, top, width, height = shape.left, shape.top, shape.width, shape.height
    shape._element.getparent().remove(shape._element)
    slide.shapes.add_picture(BytesIO(image_bytes), left, top, width, height)


def _fill_table(shape: object, rows: list[list[str]], *, body_size: float = 10.0) -> None:
    table = shape.table
    for row_index, row in enumerate(table.rows):
        values = rows[row_index] if row_index < len(rows) else [""] * len(table.columns)
        for column_index, cell in enumerate(row.cells):
            frame = cell.text_frame
            frame.clear()
            paragraph = frame.paragraphs[0]
            run = paragraph.add_run()
            run.text = str(values[column_index]) if column_index < len(values) else ""
            run.font.name = "Benton Sans BBVA"
            run.font.size = Pt(body_size + (1 if row_index == 0 else 0))
            run.font.bold = row_index == 0
            run.font.color.rgb = RGBColor(255, 255, 255) if row_index == 0 else RGBColor(7, 14, 70)
            frame.word_wrap = True
            frame.auto_size = None
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.08)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)


def _resize_table_rows(shape: object, total_height: float) -> None:
    height = Inches(total_height) // max(1, len(shape.table.rows))
    for row in shape.table.rows:
        row.height = height


def _canonical_opportunities(frame: pd.DataFrame, min_n: int) -> list[Opportunity]:
    registry = EquivalenceRegistry.default()
    raw = rank_opportunities(frame, dimensions=["Palanca"], min_n=min_n)
    seen: set[str] = set()
    output: list[Opportunity] = []
    for item in raw:
        key = registry.key("Palanca", item.value)
        if not key or key in seen:
            continue
        seen.add(key)
        output.append(item)
    return output[:3]


def _incident_records(row: pd.Series) -> list[dict[str, str]]:
    value = row.get("incident_records", [])
    if not isinstance(value, list):
        return []
    return [item for item in value if isinstance(item, dict)]


def _scenario_rows(attribution: pd.DataFrame) -> list[pd.Series]:
    if attribution is None or attribution.empty:
        return []
    ranked = attribution.copy()
    for column in ["linked_pairs", "linked_incidents", "linked_comments", "confidence"]:
        ranked[column] = pd.to_numeric(ranked.get(column, 0), errors="coerce").fillna(0)
    return [
        row
        for _, row in ranked.sort_values(["linked_pairs", "confidence"], ascending=False)
        .head(3)
        .iterrows()
    ]


def _apply_bbva_fonts(prs: Presentation) -> None:
    fonts = bbva_typography_tokens()
    for slide in prs.slides:
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    name = str(run.font.name or "").casefold()
                    run.font.name = fonts.display if "serif" in name else fonts.body


def _apply_bei_brand_lockup(prs: Presentation) -> None:
    logo_path = find_bei_logo_path()
    width = Inches(1.72)
    height = Inches(0.60)
    left = prs.slide_width - width - Inches(0.20)
    top = Inches(0.10)
    for slide in prs.slides:
        background = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        background.fill.solid()
        background.fill.fore_color.rgb = RGBColor(7, 14, 70)
        background.line.fill.background()
        slide.shapes.add_picture(
            str(logo_path),
            left + Inches(0.10),
            top + Inches(0.08),
            width=width - Inches(0.20),
            height=height - Inches(0.16),
        )


def _fit_titles_around_brand(prs: Presentation) -> None:
    """Reserve a stable title lane so dynamic copy never collides with the BEI lockup."""

    for slide in prs.slides:
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            if (
                Inches(0.25) <= shape.left <= Inches(0.50)
                and Inches(0.20) <= shape.top <= Inches(0.40)
                and shape.width >= Inches(8.0)
            ):
                shape.width = Inches(7.35)
                shape.height = Inches(0.78)
                shape.text_frame.word_wrap = True
                shape.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE


def _set_white_banner_text(shape: object, text: str) -> None:
    _set_shape_text(shape, text)
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)


def generate_exclusive_report(
    *,
    template_path: Path,
    context: ExclusiveReportContext,
    selected_nps_df: pd.DataFrame,
    comparison_nps_df: pd.DataFrame,
    attribution_df: pd.DataFrame,
    min_n: int,
) -> BusinessPptResult:
    if not template_path.exists():
        raise ValueError(f"No existe la plantilla exclusiva: {template_path}")
    prs = Presentation(str(template_path))
    metrics = _monthly_metrics(comparison_nps_df)
    if metrics.empty:
        raise ValueError("No hay una serie mensual válida para el informe exclusivo.")
    first, last = metrics.iloc[0], metrics.iloc[-1]

    slides = prs.slides
    _set_shape_text(slides[0].shapes[0], "NPS y causalidad Web")
    context_label = f"{context.service_origin} · {context.service_origin_n1} · VoC e incidencias"
    _set_shape_text(slides[0].shapes[1], context_label)

    opportunities = _canonical_opportunities(selected_nps_df, min_n=min_n)
    scenarios = _scenario_rows(attribution_df)
    lead = opportunities[0] if opportunities else None
    total_links = sum(int(float(row.get("linked_pairs", 0) or 0)) for row in scenarios)
    _set_mixed_statement(
        slides[2].shapes[4],
        "La experiencia evoluciona y mantiene fricciones",
        f"El NPS clásico pasa de {_fmt(first['nps'])} a {_fmt(last['nps'])}; la detracción cierra en {_fmt(float(last['detractor_rate']) * 100)}%.",
    )
    _set_mixed_statement(
        slides[2].shapes[6],
        f"{lead.value if lead else 'La principal palanca'} lidera el potencial",
        f"La oportunidad prioritaria concentra n={lead.n if lead else 0} y un potencial estimado de +{_fmt(lead.potential_uplift if lead else 0)} puntos.",
    )
    _set_mixed_statement(
        slides[2].shapes[5],
        f"{len(scenarios)} journeys conectan VoC e incidencias",
        f"Los escenarios priorizados reúnen {total_links} vínculos validados con trazabilidad hasta Helix.",
    )

    _set_shape_text(slides[3].shapes[7], _fmt(first["nps"]))
    _set_shape_text(slides[3].shapes[10], _fmt(last["nps"]))
    direction = "mejora" if float(last["nps"]) >= float(first["nps"]) else "retrocede"
    _set_shape_text(slides[3].shapes[1], f"El NPS {direction} y la señal queda trazada mes a mes")
    _set_shape_text(
        slides[3].shapes[3],
        f"El cierre alcanza {_fmt(last['nps'])}; la tasa detractora se sitúa en {_fmt(float(last['detractor_rate']) * 100)}%.",
    )
    _replace_picture(slides[3], 12, _chart_image(metrics))

    focus = selected_nps_df.loc[pd.to_numeric(selected_nps_df["NPS"], errors="coerce") <= 6].copy()
    topics = (
        focus.loc[
            focus["Subpalanca"].astype(str).str.strip().ne("")
            & ~focus["Subpalanca"].astype(str).str.casefold().isin({"sin comentarios", "otros"})
        ]
        .groupby("Subpalanca", dropna=False)
        .size()
        .sort_values(ascending=False)
        .head(3)
    )
    topic_rows = [["Tema", "Menciones", "Ejemplo 1", "Ejemplo 2"]]
    for topic, count in topics.items():
        examples = focus.loc[focus["Subpalanca"] == topic, "Comment"].astype(str)
        examples = [" ".join(value.split())[:48] for value in examples if value.strip()][:2]
        topic_rows.append([str(topic), str(int(count)), *(examples + [""] * (2 - len(examples)))])
    _fill_table(slides[4].shapes[7], topic_rows, body_size=9)
    _set_shape_text(
        slides[4].shapes[2],
        f"Temas más repetidos · detractores del {context.period_start:%m/%Y} al {context.period_end:%m/%Y}",
    )
    _set_shape_text(slides[4].shapes[5], "Los detractores señalan fricciones concretas y repetidas")

    analysis_df = selected_nps_df.loc[
        ~selected_nps_df["Palanca"]
        .astype(str)
        .str.casefold()
        .isin({"sin comentarios", "otros", ""})
    ].copy()
    stats = grouped_driver_stats(analysis_df, "Palanca")
    score_means = (
        analysis_df.assign(_score=pd.to_numeric(analysis_df["NPS"], errors="coerce"))
        .groupby("Palanca", dropna=False, observed=True)["_score"]
        .mean()
        .rename("mean_score")
        .reset_index()
    )
    stats = (
        stats.merge(score_means, on="Palanca", how="left").sort_values("n", ascending=False).head(7)
    )
    stat_rows = [["Palanca", "n", "Score", "% detractores"]]
    for _, row in stats.iterrows():
        stat_rows.append(
            [
                str(row["Palanca"]),
                f"{int(row['n']):,}".replace(",", "."),
                _fmt(row.get("mean_score")),
                _fmt(float(row.get("detractor_rate", 0)) * 100) + "%",
            ]
        )
    _fill_table(slides[5].shapes[7], stat_rows, body_size=10)
    _resize_table_rows(slides[5].shapes[7], 2.82)
    top_label = str(stats.iloc[0]["Palanca"]) if not stats.empty else "La principal palanca"
    _set_shape_text(slides[5].shapes[5], f"{top_label} concentra el mayor volumen observado")
    if not stats.empty:
        top = stats.iloc[0]
        _set_shape_text(
            slides[5].shapes[3],
            f"{top_label} reúne {int(top['n']):,} respuestas, score {_fmt(top.get('mean_score'))} y {_fmt(float(top.get('detractor_rate', 0)) * 100)}% de detractores. Las categorías equivalentes se muestran una sola vez.".replace(
                ",", "."
            ),
        )

    card_shapes = [(1, 3), (4, 5), (6, 7)]
    descriptions = [
        "Es la oportunidad con mayor potencial y evidencia dentro del periodo.",
        "Combina impacto estimado con una base relevante de respuestas.",
        "Completa la priorización sin duplicar etiquetas equivalentes.",
    ]
    for index, (card_index, metric_index) in enumerate(card_shapes):
        if index < len(opportunities):
            item = opportunities[index]
            _set_card(
                slides[6].shapes[card_index],
                item.value,
                descriptions[index],
            )
            _set_opportunity_metric(
                slides[6].shapes[metric_index],
                uplift=item.potential_uplift,
                confidence=item.confidence,
                n=item.n,
            )
        else:
            _set_card(
                slides[6].shapes[card_index],
                "Sin señal adicional",
                "No se fuerza una tercera oportunidad sin evidencia suficiente.",
            )
            _set_shape_text(slides[6].shapes[metric_index], "n/d")

    journey_rows = [["Journey", "Touchpoint", "Palanca", "Subpalanca", "Tópico NPS ancla"]]
    for row in scenarios:
        journey_rows.append(
            [
                str(row.get("touchpoint", row.get("nps_topic", "Journey"))),
                str(row.get("journey_route", row.get("touchpoint", "")))[:52],
                str(row.get("palanca", "")),
                str(row.get("subpalanca", "")),
                str(row.get("nps_topic", "")),
            ]
        )
    _fill_table(slides[8].shapes[7], journey_rows, body_size=9)
    _resize_table_rows(slides[8].shapes[7], 2.62)
    slides[8].shapes[4].top = Inches(5.13)
    _set_shape_text(
        slides[8].shapes[5], f"{len(scenarios)} journeys reúnen {total_links} vínculos validados"
    )

    if scenarios:
        first_scenario = scenarios[0]
        records = _incident_records(first_scenario)[:5]
        frame = slides[9].shapes[1].text_frame
        frame.clear()
        paragraph = frame.paragraphs[0]
        run = paragraph.add_run()
        run.text = (
            f"{first_scenario.get('touchpoint', first_scenario.get('nps_topic', 'Escenario prioritario'))}: "
            f"{int(first_scenario.get('linked_pairs', 0) or 0)} vínculos, "
            f"{int(first_scenario.get('linked_incidents', 0) or 0)} incidencias y "
            f"{int(first_scenario.get('linked_comments', 0) or 0)} comentarios.\n\n"
            "Referencias Helix: "
        )
        for index, record in enumerate(records):
            incident_id = str(record.get("incident_id", record.get("id", ""))).strip()
            if not incident_id:
                continue
            link = frame.paragraphs[-1].add_run()
            link.text = (", " if index else "") + incident_id
            link.hyperlink.address = str(
                record.get("url", "") or context.helix_base_url + incident_id
            )
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE

    for slide_index, scenario_index in [(10, 1), (11, 2)]:
        if scenario_index >= len(scenarios):
            continue
        row = scenarios[scenario_index]
        title = str(row.get("touchpoint", row.get("nps_topic", "Journey priorizado")))
        if slide_index == 10:
            _set_shape_text(
                slides[slide_index].shapes[5], f"{title} concentra evidencia convergente"
            )
            _set_shape_text(
                slides[slide_index].shapes[1],
                f"{int(row.get('linked_pairs', 0) or 0)} vínculos validados",
            )
            _set_shape_text(
                slides[slide_index].shapes[4],
                f"{int(row.get('linked_incidents', 0) or 0)} incidencias Helix y {int(row.get('linked_comments', 0) or 0)} comentarios sustentan la hipótesis.",
            )
            _set_shape_text(
                slides[slide_index].shapes[8], f"Confianza {_fmt(row.get('confidence', 0))}"
            )
        else:
            probability = pd.to_numeric(
                pd.Series([row.get("detractor_probability")]), errors="coerce"
            ).iloc[0]
            probability_label = (
                f"{_fmt(float(probability) * 100)}%" if pd.notna(probability) else "n/d"
            )
            _set_shape_text(
                slides[slide_index].shapes[11], f"{title} completa el mapa causal prioritario"
            )
            _set_white_banner_text(
                slides[slide_index].shapes[2],
                f"{int(row.get('linked_pairs', 0) or 0)} vínculos validados",
            )
            _set_white_banner_text(
                slides[slide_index].shapes[6],
                f"Detracción estimada · {probability_label}",
            )
            _set_shape_text(
                slides[slide_index].shapes[5],
                (
                    f"El escenario presenta una probabilidad de detracción del {probability_label} "
                    f"y una confianza de {_fmt(row.get('confidence', 0), 2)}."
                    if pd.notna(probability)
                    else "La evidencia converge, pero no permite estimar la detracción con robustez; "
                    "el escenario se mantiene como hipótesis a validar."
                ),
            )
            _set_shape_text(
                slides[slide_index].shapes[1],
                f"{int(row.get('linked_incidents', 0) or 0)} incidencias y {int(row.get('linked_comments', 0) or 0)} comentarios convergen en este journey.",
            )

    _fit_titles_around_brand(prs)
    _apply_bbva_fonts(prs)
    _apply_bei_brand_lockup(prs)
    output = BytesIO()
    prs.save(output)
    file_name = (
        f"nps-lens-exclusivo-{context.service_origin_n1.casefold().replace(' ', '-')}-"
        f"{context.period_end.strftime('%Y%m%d')}.pptx"
    )
    return BusinessPptResult(
        file_name=file_name, content=output.getvalue(), slide_count=len(slides)
    )
