"""Branding stays readable, single-sourced and outside the analytical content."""

import json
import subprocess
from base64 import b64decode
from datetime import date
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from test_business_ppt import _sample_payload

from nps_lens.design.brand import BRAND, BRAND_ASSETS, EMAIL_SIGNATURE, email_signature
from nps_lens.platform.publication import _newsletter
from nps_lens.platform.webapp_preview import build_preview, load_publication
from nps_lens.reports.executive_ppt import generate_business_review_ppt


@pytest.fixture(scope="module")
def branded_decks():
    payload = _sample_payload()
    result = generate_business_review_ppt(
        service_origin="México",
        service_origin_n1="",
        service_origin_n2="",
        period_start=date(2026, 1, 1),
        period_end=date(2026, 1, 31),
        focus_name="detractores",
        attribution_df=payload["attribution"],
        selected_nps_df=payload["selected_nps"],
        comparison_nps_df=payload["comparison_nps"],
    )
    return result


def test_branding_signs_every_slide_with_unique_numbering_and_reuses_artwork(branded_decks):
    for content in (branded_decks.content, branded_decks.compact_content):
        deck = Presentation(BytesIO(content))
        assert BRAND["name"] in deck.core_properties.author
        for index, slide in enumerate(deck.slides):
            names = [shape.name for shape in slide.shapes]
            assert names.count("Report page number") == names.count("bIA") == 1
            assert names.count(BRAND["name"]) == 1
            page = next(shape for shape in slide.shapes if shape.name == "Report page number")
            assert page.text == str(index + 1)
            footer = next(shape for shape in slide.shapes if shape.name == "Corporate scope footer")
            assert BRAND["name"] in footer.text and "VoC:" in footer.text
            for shape in slide.shapes:
                assert shape.left + shape.width <= deck.slide_width
                assert shape.top + shape.height <= deck.slide_height
            if not index:
                corporate = next(shape for shape in slide.shapes if shape.name == BRAND["name"])
                initiative = next(shape for shape in slide.shapes if shape.name == "bIA")
                title, scope = slide.shapes[0], slide.shapes[1]
                assert corporate.left == Inches(0.38)
                assert initiative.left > Inches(8)
                assert corporate.top == initiative.top == Inches(0.25)
                assert scope.top == title.top + title.height + Inches(0.08)
            if index:
                callout = next(shape for shape in slide.shapes if shape.name == "Report conclusion")
                title = next(shape for shape in slide.shapes if shape.name == "Report title")
                corporate = next(shape for shape in slide.shapes if shape.name == BRAND["name"])
                logo = next(shape for shape in slide.shapes if shape.name == "bIA")
                assert title.left + title.width < corporate.left
                assert title.top == corporate.top
                assert title.top + title.height <= Inches(1.30)
                assert callout.top + callout.height <= logo.top
        for master in deck.slide_masters:
            for template in (master, *master.slide_layouts):
                assert not any(
                    shape._element.xpath(".//a:fld[@type='slidenum']") for shape in template.shapes
                )
                assert not any(
                    ph.get("type") == "sldNum"
                    for shape in template.shapes
                    for ph in shape._element.xpath(".//p:ph")
                )
        with ZipFile(BytesIO(content)) as archive:
            images = [
                archive.read(name) for name in archive.namelist() if name.startswith("ppt/media/")
            ]
        assert images.count((BRAND_ASSETS / "bia.png").read_bytes()) == 1
    history = Presentation(BytesIO(branded_decks.content)).slides[1]
    explanation = next(shape for shape in history.shapes if shape.name == "Historical explanation")
    assert explanation.top + explanation.height <= Inches(4.90)
    assert len(explanation.text_frame.paragraphs) == 3
    assert "NPS clásico de la base histórica" in explanation.text


def test_newsletter_and_browser_assets_share_the_signature_and_font(tmp_path):
    preview = build_preview(
        Path("webapp/apps-script"), tmp_path, load_publication(None)
    ).read_text()
    assert "<?= BRAND." not in preview
    assert BRAND["initiative_name"] in preview
    assert 'format("woff2")' in preview
    for weight in ("Book", "Medium", "Bold"):
        font = Path(f"frontend/public/assets/fonts/bbva/BentonSansBBVA-{weight}.woff2")
        assert font.read_bytes().startswith(b"wOF2")
    html = _newsletter({"newsletter": {"brand": BRAND["name"]}}, "report.pptx").decode()
    assert EMAIL_SIGNATURE in html
    script = Path("webapp/apps-script/00_Brand.gs").read_text()
    script += "\nprocess.stdout.write(JSON.stringify(BRAND));"
    generated = json.loads(
        subprocess.run(["node", "-e", script], capture_output=True, check=True, text=True).stdout
    )
    assert generated["newsletter_prefix"] == "[bIA]"
    assert generated["email_signature"] == email_signature("cid:bia-logo")
    assert b64decode(generated["initiative_logo"]) == (BRAND_ASSETS / "bia.png").read_bytes()
    with Image.open(BRAND_ASSETS / "bia.png") as logo:
        assert logo.size == (852, 520)
    assert 'viewBox="0 0 213 130"' in Path("frontend/public/assets/brand/bia.svg").read_text()
    assert "data:image/png;base64," in EMAIL_SIGNATURE
