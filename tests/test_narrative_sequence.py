from datetime import date
from io import BytesIO

import pandas as pd
import pytest
from pptx import Presentation

from nps_lens.analytics.causal_evidence import engine_quality, link_confidence_label
from nps_lens.reports.executive_newsletter import build_executive_newsletter
from nps_lens.reports.executive_ppt import generate_business_review_ppt
from nps_lens.reports.narrative import comment_groups, experience_topics, ordered_scenarios
from nps_lens.services.dashboard_service import DashboardService


def opinions():
    return pd.DataFrame(
        {
            "ID": range(8),
            "Palanca": ["Operativa"] * 3 + ["Continuidad"] * 3 + ["Acceso"] * 2,
            "Subpalanca": ["Fallo"] * 8,
            "NPS": [0, 3, 3, 2, 4, 6, None, None],
            "Comment": ["No puedo completar mi operación"] * 8,
            "Canal": ["Web"] * 8,
            "Fecha": ["2026-08-01"] * 8,
        }
    )


def cases():
    rows = []
    for topic, value, identity, volume, text in [
        ("Continuidad", 0, "c", 1, "La aplicación no funciona"),
        ("Operativa", 2.5, "b", 2, "No puedo consultar movimientos"),
        ("Operativa", 2, "a", 2, "No pude pagar sueldos"),
        ("Operativa", None, "d", 100, "No puedo hacer una operación"),
        ("Acceso", 0, "e", 1, "No puedo iniciar sesión"),
    ]:
        rows.append(
            {
                "scenario_id": identity,
                "anchor_topic": topic + " > Fallo",
                "nps_topic": topic + " > Fallo",
                "avg_nps": value,
                "linked_comments": volume,
                "linked_incidents": 1,
                "linked_pairs": volume,
                "causal_engine": "llm",
                "avg_semantic_confidence": 0.9,
                "affected_task": "operar",
                "observed_symptom": "la operación falla",
                "comment_records": [{"comment_id": identity, "nps": value, "comment": text}],
                "incident_records": [
                    {
                        "incident_id": "INC" + identity,
                        "summary": "La operación falla",
                        "url": "https://helix.example/" + identity,
                    }
                ],
            }
        )
    return pd.DataFrame(rows)


def test_order_uses_topic_population_then_linked_score_and_preserves_every_case():
    topics = experience_topics(opinions(), channel="Web")
    assert topics.value.tolist() == ["Operativa", "Continuidad", "Acceso"]
    assert topics.score.iloc[:2].tolist() == [2, 4]
    assert pd.isna(topics.score.iloc[2])
    expected = ["a", "b", "d", "c", "e"]
    assert ordered_scenarios(cases(), topics).scenario_id.tolist() == expected
    assert (
        ordered_scenarios(cases().sample(frac=1, random_state=1), topics).scenario_id.tolist()
        == expected
    )
    assert len(ordered_scenarios(cases(), topics)) == len(cases())
    assert experience_topics(opinions(), channel="App").empty


def test_ties_use_volume_then_normalized_name_and_stable_identity():
    frame = opinions()
    frame.loc[:, "NPS"] = 2
    topics = experience_topics(frame)
    assert topics.value.tolist() == ["Continuidad", "Operativa", "Acceso"]
    rows = cases().iloc[:3].copy()
    rows.loc[:, "anchor_topic"] = "Operativa > Fallo"
    rows.loc[:, "avg_nps"] = 0
    assert ordered_scenarios(rows, topics).scenario_id.tolist() == ["a", "b", "c"]


def test_invalid_scenario_average_sorts_after_valid_zero():
    rows = cases().iloc[:3].copy()
    rows.loc[:, "anchor_topic"] = "Operativa > Fallo"
    rows.loc[:, "avg_nps"] = [99, 0, -1]
    assert ordered_scenarios(rows, experience_topics(opinions())).scenario_id.tolist() == [
        "b",
        "a",
        "c",
    ]


def test_scenario_card_keeps_anchor_only_in_the_descriptive_fact_sheet():
    cards = object.__new__(DashboardService)._build_linking_scenario_cards(cases())
    for card in cards:
        assert card["title"] == "operar → la operación falla"
        assert card["identity_rows"][0] == {
            "label": "Tópico NPS ancla",
            "value": card["anchor_topic"],
        }
        assert card["anchor_topic"] not in card["title"]


def test_scenario_score_labels_keep_full_aggregate_counts_with_partial_records():
    source = cases().iloc[:1].copy()
    record = {"comment_id": "c", "nps": 0, "comment": "No puedo operar"}
    source.at[0, "comment_records"] = [record]
    source["score_distribution"] = [[{"score": 0, "count": 99, "label": "Etiqueta obsoleta"}]]
    card = object.__new__(DashboardService)._build_linking_scenario_cards(source)[0]
    assert card["score_distribution"] == [
        {"score": 0, "count": 99, "label": "Score 0 / 99 Comentarios"}
    ]


def test_webapp_renders_anchor_once_in_fact_sheet_and_uses_the_short_heading():
    import json
    import subprocess
    from pathlib import Path

    source = Path("webapp/apps-script/App.html").read_text()
    renderer = source[
        source.index("  function scenarioCard(") : source.index("  function scenarioWorkspace(")
    ]
    card = object.__new__(DashboardService)._build_linking_scenario_cards(cases())[0]
    script = (
        """
const assert=require('node:assert/strict');
const state={scenarioDetail:'helix',evidenceView:'table'};
const rows=value=>Array.isArray(value)?value:[];
const esc=value=>String(value);
const scenarioEvidence=()=>'';
const contentTabs=()=>'';
"""
        + renderer
        + "\nconst card="
        + json.dumps(card)
        + ";\n"
        + """
const html=scenarioCard(card,0,1);
assert.ok(html.includes('<h3>operar → la operación falla</h3>'));
assert.equal(html.split(card.anchor_topic).length-1,1);
assert.ok(html.includes('<dt>Tópico NPS ancla</dt><dd>'+card.anchor_topic+'</dd>'));
"""
    )
    subprocess.run(["node"], input=script, check=True, capture_output=True, text=True)


def test_groups_count_unique_comments_sort_scores_and_preserve_zero_and_missing():
    records = [
        {"comment_id": "a", "nps": 3, "comment": "No puedo operar"},
        {"comment_id": "b", "nps": 0, "comment": "Error"},
        {"comment_id": "c", "nps": 0, "comment": "Error"},
        {"comment_id": "d", "nps": None, "comment": "Sin nota"},
    ]
    groups = comment_groups(records + [records[1]])
    assert [(g["score"], g["count"]) for g in groups] == [(0, 2), (3, 1), (None, 1)]
    assert [g["label"] for g in groups] == [
        "Score 0 / 2 Comentarios",
        "Score 3 / 1 Comentario",
        "Sin score / 1 Comentario",
    ]
    assert sum(g["count"] for g in groups) == 4


@pytest.mark.parametrize(
    "engine,column,label",
    [
        ("llm", "avg_semantic_confidence", "CONFIANZA SEMÁNTICA"),
        ("rules", "avg_text_similarity", "SIMILITUD TEXTUAL"),
    ],
)
def test_quality_uses_own_engine_scale_and_distinguishes_zero(engine, column, label):
    assert link_confidence_label(engine) == label
    assert engine_quality({"causal_engine": engine, column: 0}) == 0
    assert engine_quality({"causal_engine": engine, column: 0.9}) == 0.9
    assert engine_quality({"causal_engine": engine, column: 90}) is None
    assert engine_quality({"causal_engine": engine}) is None


def test_newsletter_takes_one_traceable_quote_per_first_two_topics():
    frame = opinions()
    cards = object.__new__(DashboardService)._build_linking_scenario_cards(
        ordered_scenarios(cases(), experience_topics(frame))
    )
    result = build_executive_newsletter(
        current_df=frame,
        period_kpis={},
        linking={"scenarios": {"cards": cards}},
        topic_channel="Todos",
        period_start=date(2026, 8, 1),
        period_end=date(2026, 8, 31),
    )
    assert result["quotes"] == ["No pude pagar sueldos", "La aplicación no funciona"]
    assert [r["label"] for r in result["signals"]] == ["Operativa", "Continuidad", "Acceso"]
    assert [r["scenario_id"] for r in cards] == ["a", "b", "d", "c", "e"]
    single = {"scenarios": {"cards": cards[:1]}}
    result = build_executive_newsletter(
        current_df=frame,
        period_kpis={},
        linking=single,
        topic_channel="Todos",
        period_start=date(2026, 8, 1),
        period_end=date(2026, 8, 31),
    )
    assert result["quotes"] == ["No pude pagar sueldos"]


def test_ppt_interleaves_topics_and_cases_in_the_same_sequence():
    frame = opinions()
    out = generate_business_review_ppt(
        service_origin="Fixture",
        service_origin_n1="",
        service_origin_n2="",
        period_start=date(2026, 8, 1),
        period_end=date(2026, 8, 31),
        focus_name="Todos",
        selected_nps_df=frame,
        comparison_nps_df=frame,
        attribution_df=cases(),
    )
    deck = Presentation(BytesIO(out.compact_content))
    assert len(deck.slides) == 4 + 3 + 5
    sequence = []
    for slide in list(deck.slides)[4:]:
        separator = next(
            (shape.text for shape in slide.shapes if shape.name == "Topic separator"), None
        )
        if separator:
            sequence.append(separator)
        else:
            values = [shape.text for shape in slide.shapes if shape.name == "Scenario metric value"]
            sequence.append(values[0])
    assert sequence == [
        "VoC : Operativa",
        "2,00",
        "2,50",
        "n/d",
        "VoC : Continuidad",
        "0,00",
        "VoC : Acceso",
        "0,00",
    ]


def test_webapp_navigation_uses_the_published_sequence_and_resets_on_new_publication():
    import json
    import subprocess
    from pathlib import Path

    source = Path("webapp/apps-script/App.html").read_text()
    workspace = source[
        source.index("  function scenarioWorkspace(") : source.index(
            "  function linkingDiagnostics("
        )
    ]
    apply_start = source.index("  function applyPublication(")
    assert (
        "state.scenarioIndex=0;"
        in source[apply_start : source.index("  function ", apply_start + 10)]
    )
    rows = ordered_scenarios(cases(), experience_topics(opinions())).to_dict("records")
    ids = [row["scenario_id"] for row in rows]
    script = (
        """
const assert=require('node:assert/strict');
const state={scenarioIndex:0};
const rows=value=>value;
const esc=value=>String(value);
const scenarioCard=card=>card.scenario_id;
"""
        + workspace
        + "\nconst cards="
        + json.dumps(rows)
        + ";\nconst expected="
        + json.dumps(ids)
        + ";\n"
        + """
for(let i=0;i<cards.length;i++){
  state.scenarioIndex=i;
  assert.ok(scenarioWorkspace(cards).endsWith(expected[i]));
}
state.scenarioIndex=500;
scenarioWorkspace(cards);
assert.ok(state.scenarioIndex>=0&&state.scenarioIndex<cards.length);
"""
    )
    subprocess.run(["node"], input=script, check=True, capture_output=True, text=True)


def test_ppt_has_one_top_four_overview_and_keeps_evidence_outside_that_top():
    frame = pd.DataFrame(
        {
            "Palanca": [f"Topic {i}" for i in range(6)],
            "Subpalanca": ["Fallo"] * 6,
            "Comment": ["No puedo completar la operación"] * 6,
            "NPS": list(range(6)),
            "Fecha": ["2026-08-01"] * 6,
        }
    )
    chains = cases().iloc[:2].copy()
    chains["anchor_topic"] = ["Topic 4 > Fallo", "Topic 5 > Fallo"]
    chains["nps_topic"] = chains.anchor_topic
    report = generate_business_review_ppt(
        service_origin="Fixture",
        service_origin_n1="",
        service_origin_n2="",
        period_start=date(2026, 8, 1),
        period_end=date(2026, 8, 31),
        focus_name="Todos",
        selected_nps_df=frame,
        comparison_nps_df=frame,
        attribution_df=chains,
    )
    deck = Presentation(BytesIO(report.compact_content))
    overview = [s for s in deck.slides if any(sh.name == "Topic heading" for sh in s.shapes)]
    assert len(overview) == 1
    assert [sh.text for sh in overview[0].shapes if sh.name == "Topic heading"] == [
        f"Topic {i}" for i in range(4)
    ]
    title = " ".join(
        next(sh.text for sh in overview[0].shapes if sh.name == "Report title").split()
    )
    assert "top 4 de los tópicos" in title and "(1/" not in title
    assert [sh.text for s in deck.slides for sh in s.shapes if sh.name == "Topic separator"] == [
        "VoC : Topic 4",
        "VoC : Topic 5",
    ]


def test_quotes_do_not_silently_replace_a_duplicate_second_topic_with_a_third():
    rows = cases()
    rows.at[0, "comment_records"] = [
        {"comment_id": "c", "nps": 0, "comment": "No pude pagar sueldos"}
    ]
    result = build_executive_newsletter(
        current_df=opinions(),
        period_kpis={},
        linking={"scenarios": {"cards": rows.to_dict("records")}},
        topic_channel="Todos",
        period_start=date(2026, 8, 1),
        period_end=date(2026, 8, 31),
    )
    assert result["quotes"] == ["No pude pagar sueldos"]


def test_null_comment_records_and_unknown_topic_do_not_become_zero():
    rows = pd.DataFrame(
        [
            {
                "nps_topic": "Acceso > Fallo",
                "anchor_topic": float("nan"),
                "palanca": "Acceso",
                "avg_nps": None,
                "comment_records": float("nan"),
            }
        ]
    )
    result = ordered_scenarios(rows, experience_topics(opinions()))
    assert result.narrative_topic.tolist() == ["Acceso"]
    assert result.avg_nps.isna().all()
    assert comment_groups([{"comment_id": "x", "nps": 2.5, "comment": "Error"}])[0]["score"] is None


@pytest.mark.parametrize("incident_count", [3, 40])
def test_ppt_case_is_one_slide_with_detailed_first_incident_and_overflow_ids(incident_count):
    from nps_lens.reports.evidence_layout import EVIDENCE_LAYOUT

    rows = cases().iloc[:1].copy()
    records = [
        {
            "incident_id": f"INC000104{i:06d}",
            "summary": "La aplicación queda conectando y no permite realizar la operación " * 4,
            "url": f"https://helix.example/{i}",
        }
        for i in range(incident_count)
    ]
    rows.at[0, "incident_records"] = records
    report = generate_business_review_ppt(
        service_origin="Fixture",
        service_origin_n1="",
        service_origin_n2="",
        period_start=date(2026, 8, 1),
        period_end=date(2026, 8, 31),
        focus_name="Todos",
        selected_nps_df=opinions(),
        comparison_nps_df=opinions(),
        attribution_df=rows,
    )
    deck = Presentation(BytesIO(report.compact_content))
    slides = [s for s in deck.slides if any(sh.name == "Incident evidence" for sh in s.shapes)]
    assert len(slides) == 1
    slide = slides[0]
    title = next(sh.text for sh in slide.shapes if sh.name == "Report title")
    assert "(1/" not in title
    cells = [sh for sh in slide.shapes if sh.name == "Incident evidence"]
    assert records[0]["incident_id"] in cells[0].text
    assert "La aplicación queda conectando" in cells[0].text
    assert f"{incident_count - 1} Incidencias:" in cells[-1].text
    assert f"[{records[1]['incident_id']}]" in cells[-1].text
    assert ("…" in cells[-1].text) == (incident_count == 40)
    assert all((sh.top + sh.height) / 914400 <= EVIDENCE_LAYOUT.body_bottom + 0.001 for sh in cells)
    for record in records:
        assert record["incident_id"] in slide.notes_slide.notes_text_frame.text
    assert any(
        run.hyperlink.address == records[0]["url"]
        for paragraph in cells[0].text_frame.paragraphs
        for run in paragraph.runs
    )
