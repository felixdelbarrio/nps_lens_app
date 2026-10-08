"""Cross-channel evidence contracts, deterministic and isolated from user data/LLMs."""

from __future__ import annotations

import copy
import json
from datetime import date
from io import BytesIO

import pandas as pd
import pytest
from pptx import Presentation
from test_business_ppt import _sample_payload
from test_taxonomy_exchange import (
    classifier_files,
    classifier_zip,
    designer_zip,
    exported,
)
from test_taxonomy_exchange import exchange_fixture as exchange_fixture

from nps_lens.analytics.causal_evidence import (
    CausalEvidenceEvaluator,
    link_confidence_label,
    scenario_impact_score,
)
from nps_lens.analytics.drivers import driver_table
from nps_lens.analytics.incident_attribution import build_incident_attribution_chains
from nps_lens.analytics.nps_helix_link import build_incident_text, nps_matchable_mask
from nps_lens.analytics.signal_quality import audit_classifications, signal_quality
from nps_lens.domain.privacy import redact_operational_snippet, redact_public_payload
from nps_lens.reports import executive_ppt
from nps_lens.reports.coherence import (
    ReportCoherenceError,
    validate_classification_context,
    validate_metric_payload,
)
from nps_lens.reports.content_selectors import select_causal_scenarios, select_negative_delta_rows
from nps_lens.reports.executive_newsletter import build_executive_newsletter
from nps_lens.services.analytics.kpis_service import build_period_kpis
from nps_lens.services.dashboard_service import DashboardService
from nps_lens.services.helix_exchange import EvidenceLink
from nps_lens.services.taxonomy_prompts import PROJECT_INSTRUCTIONS


def link(**overrides):
    return (
        dict(
            nps_id="n1",
            incident_id="i1",
            comment_quote="No genera token",
            incident_quote="Error al generar token",
            comment_date="2026-09-03",
            incident_start_date="2026-09-01",
            nps_score=2,
            nps_group="DETRACTOR",
            palanca="Acceso",
            subpalanca="Token no generado",
            semantic_confidence=0.9,
            causal_engine="llm",
            same_task=True,
            same_symptom=True,
            linked_comments=3,
            linked_incidents=2,
            avg_score=2,
            detractor_rate=1,
            **{},
        )
        | overrides
    )


@pytest.mark.parametrize(
    ("changes", "level"),
    [
        ({}, "EVIDENCIA_OPERATIVA_RECURRENTE"),
        ({"linked_incidents": 1}, "ASOCIACION_OPERATIVA_CONSISTENTE"),
        ({"avg_score": 9}, "ASOCIACION_OPERATIVA_CONSISTENTE"),
        ({"detractor_rate": 0.5}, "ASOCIACION_OPERATIVA_CONSISTENTE"),
        ({"same_task": False}, "INDICIO_SEMANTICO"),
        ({"same_symptom": False}, "INDICIO_SEMANTICO"),
        ({"comment_quote": ""}, "INDICIO_SEMANTICO"),
        ({"comment_date": None}, "INDICIO_SEMANTICO"),
        ({"comment_date": "2026-08-31"}, "INDICIO_SEMANTICO"),
        ({"comment_date": "2027-01-01"}, "INDICIO_SEMANTICO"),
        ({"incident_close_date": "2026-08-01"}, "INDICIO_SEMANTICO"),
        ({"semantic_confidence": 0.1}, "CAUSALIDAD_NO_ACREDITADA"),
        ({"incident_id": ""}, "SIN_EVIDENCIA"),
        ({"subpalanca": "Información insuficiente"}, "SIN_EVIDENCIA"),
    ],
)
def test_evidence_requires_all_defensible_conditions(changes, level):
    result = CausalEvidenceEvaluator().evaluate(link(**changes))
    assert result["evidence_level"] == level
    assert "causalidad demostrada" not in result["evidence_reason"].lower()


def test_duplicate_pairs_do_not_inflate_recurrence_and_unknown_pairs_do_not_inherit_strength():
    row = link(
        nps_score=2,
        nps_date="2026-09-03",
        incident_date="2026-09-01",
        incident_summary="token falla",
    )
    rows = pd.DataFrame([row] * 10)
    assert (
        CausalEvidenceEvaluator().evaluate_scenario(rows)["evidence_level"]
        == "ASOCIACION_OPERATIVA_CONSISTENTE"
    )
    rows = pd.DataFrame([row, row | {"nps_id": "n2", "incident_id": "i2", "same_task": False}])
    assert (
        CausalEvidenceEvaluator().evaluate_scenario(rows)["evidence_level"] == "INDICIO_SEMANTICO"
    )


def test_reserves_never_lead_insights_but_remain_in_quality_and_kpis():
    frame = pd.DataFrame(
        {
            "Palanca": ["Sin clasificación temática"] * 10 + ["Acceso"],
            "Subpalanca": ["Información insuficiente"] * 10 + ["Token"],
            "Comment": ["NO"] * 10 + ["No genera token"],
            "NPS": [0] * 10 + [2],
        }
    )
    assert {r.value for r in driver_table(frame, "Subpalanca")} == {
        "Información insuficiente",
        "Token",
    }
    assert nps_matchable_mask(frame).sum() == 1
    quality = signal_quality(frame)
    assert quality["insufficient_comments"] == 10 and quality["warnings"]
    delta = pd.DataFrame(
        {
            "value": ["Información insuficiente", "Token"],
            "delta_nps": [-100, -20],
            "n_current": [100, 5],
        }
    )
    assert select_negative_delta_rows(delta, max_rows=10).value.tolist() == ["Token"]
    assert executive_ppt._period_overview(frame)["friction"]["topic"] == "Acceso > Token"


def test_ranking_prioritizes_unique_incidents_before_semantic_quality():
    strong = dict(
        nps_topic="Token",
        linked_comments=8,
        linked_incidents=4,
        avg_nps=2,
        detractor_rate=1,
        avg_text_similarity=0.9,
        quote_coverage=1,
        narrative_coverage=1,
        freshness=0.9,
    )
    weak = strong | dict(
        nps_topic="Lentitud",
        linked_comments=1000,
        linked_incidents=40,
        avg_nps=9,
        detractor_rate=0.1,
        avg_text_similarity=0.35,
        quote_coverage=0,
    )
    reserve = strong | dict(nps_topic="Sin clasificación temática > Información insuficiente")
    assert scenario_impact_score(strong) > scenario_impact_score(weak)
    result = select_causal_scenarios(pd.DataFrame([weak, reserve, strong]), max_rows=10)
    assert result.nps_topic.tolist() == ["Lentitud", "Token"]
    assert select_causal_scenarios(pd.DataFrame([reserve]), max_rows=10).empty


def test_short_comments_are_reviewed_without_keyword_reclassification():
    comments = {
        "a": "No funciona bien la página",
        "b": "Muy lenta",
        "c": "No funcionó el token",
        "d": "PÉSIMO",
        "e": "NO",
    }
    assignments = {
        key: {"primary_classification": {"sublever": "Información insuficiente"}}
        for key in comments
    }
    before = copy.deepcopy(assignments)
    audit = audit_classifications(comments, assignments)
    assert audit["suspicious_ids"] == ["a", "b", "c", "d"] and audit["review_required"]
    assert assignments == before
    assert "No funciona bien la página" in PROJECT_INSTRUCTIONS["classifier"]


def test_suspicious_import_reports_without_heuristic_quality_gate(exchange):
    handler, ctx, frame, _ = exchange
    frame.loc[:, "Comment"] = "No funciona bien la página"
    handler.import_response(ctx, designer_zip(handler, ctx), "designer")
    handler.taxonomy.configure(ctx, {"accept_proposal": True})
    handler.taxonomy.configure(ctx, {"active": "DISCOVERED"})
    request = exported(handler.export(ctx, "classifier")["saved_paths"])
    response = classifier_files(request["manifest.json"], request)
    reserve = next(
        k
        for k, c in request["taxonomy.json"]["categories"].items()
        if c["sublever"] == "Información insuficiente"
    )
    for batch in response["results"].values():
        for row in batch["classifications"]:
            row["primary"] = reserve
    result = handler.import_response(ctx, classifier_zip(response), "classifier")
    assert result["classification_audit"]["suspicious_rate"] == 1
    assert handler.progress(ctx)["received"] == len(frame)


def test_incident_semantic_contract_prioritizes_cheque_description_over_administrative_token():
    narrative = "No tiene opción de depósito de cheque físico"
    frame = pd.DataFrame(
        [
            {
                "Incident Number": "I1",
                "Summary": "Acceso / Claves / Token",
                "Detailed Description": narrative,
            }
        ]
    )
    assert narrative in build_incident_text(frame).iloc[0]
    instructions = PROJECT_INSTRUCTIONS["helix"]
    assert "routing y campos de plantilla no son evidencia" in instructions
    assert "clasifica por cheques" in instructions
    # Semantic decisions are supplied externally; no production keyword classifier.
    validated = EvidenceLink(
        nps_id="N1",
        confidence=0.9,
        incident_quote=narrative,
        comment_quote="No permite depositar cheque físico",
        same_task=True,
        same_symptom=True,
        affected_task="Depositar cheque físico",
        observed_symptom="Opción no disponible",
    )
    assert validated.affected_task == "Depositar cheque físico"


PII = "CUIT: 30-12345678-9; Razón social: Acme SA; Cliente: Juan Pérez; email: juan@example.com; teléfono: +54 11 4444 5555; código de empresa: ABC999; No genera token."


def test_redaction_is_shared_preserves_operational_signal_and_does_not_leak_split_spans():
    redacted = redact_operational_snippet(PII)
    for secret in ("12345678", "Acme", "Juan Pérez", "juan@example", "4444", "ABC999"):
        assert secret not in redacted
    assert "No genera token" in redacted
    payload = {
        "comment": PII,
        "comment_id": "n1",
        "url": "https://example.com/I1",
        "comment_segments": [
            {"text": "juan@", "bold": True},
            {"text": "example.com", "bold": False},
        ],
    }
    safe = redact_public_payload(payload)
    assert safe["comment_id"] == "n1" and safe["url"] == payload["url"]
    assert "juan@" not in json.dumps(safe)
    assert (
        DashboardService._serialize_rows(pd.DataFrame([{"Comment": PII}]))[0]["Comment"] == redacted
    )


def test_metric_corruption_fails_and_monthly_headline_matches_printed_table():
    df = pd.DataFrame(
        {
            "Fecha": pd.to_datetime(
                ["2026-08-02", "2026-08-03", "2026-09-01", "2026-09-02", "2026-09-03"]
            ),
            "NPS": [9, 10, 0, 9, 9],
            "Comment": ["ok"] * 5,
        }
    )
    current = df.iloc[2:]
    kpis = build_period_kpis(
        history_df=df,
        current_df=current,
        pop_year="2026",
        pop_month="09",
        context_label="Septiembre",
    )
    overview = executive_ppt._period_overview(current, period_kpis=kpis)
    assert overview["classic_delta"] == kpis["period"]["deltas"]["classic_nps"]["value"]
    assert (
        overview["classic_delta"] < 0 < kpis["period"]["temporal"]["deltas"]["classic_nps"]["value"]
    )
    for metric in ("classic_nps", "nps_average", "promoter_rate", "detractor_rate", "comments"):
        corrupt = copy.deepcopy(kpis)
        corrupt["period"]["deltas"][metric]["value"] = 12345
        with pytest.raises(ReportCoherenceError):
            validate_metric_payload(corrupt)
    corrupt = copy.deepcopy(kpis)
    corrupt["period"]["deltas"]["classic_nps"]["display"] = "+99,00 pts"
    with pytest.raises(ReportCoherenceError):
        validate_metric_payload(corrupt)


@pytest.mark.parametrize("field", ["taxonomy_fingerprint", "scope", "causal_engine"])
def test_report_preflight_rejects_incompatible_artifacts(field):
    scope = {"company": "Bank", "channel": "Web", "year": "2026", "month": "09"}
    artifact = {"taxonomy_fingerprint": "fp", "scope": scope, "causal_engine": "llm"}
    validate_classification_context(
        "fp", {"comentarios": artifact}, expected_scope=scope, causal_engine="llm"
    )
    with pytest.raises(ReportCoherenceError, match="regenera"):
        validate_classification_context(
            "fp",
            {"comentarios": artifact | {field: "wrong"}},
            expected_scope=scope,
            causal_engine="llm",
        )


def test_ppt_newsletter_and_dashboard_share_confidence_copy_and_safe_evidence(monkeypatch):
    monkeypatch.setattr(executive_ppt, "_figure_png", lambda *a, **kw: None)
    payload = _sample_payload()
    attribution = payload["attribution"].copy()
    attribution["causal_engine"] = "llm"
    attribution["evidence_reason"] = "Asociación operativa consistente."
    attribution.at[0, "incident_examples"] = [PII]
    attribution.at[0, "comment_examples"] = [PII]
    attribution.at[0, "comment_records"] = [{"comment_id": "n1", "comment": PII}]
    attribution.at[0, "incident_records"] = [{"incident_id": "I1", "summary": PII}]
    report = executive_ppt.generate_business_review_ppt(
        service_origin="Bank",
        service_origin_n1="",
        service_origin_n2="",
        period_start=date(2026, 1, 1),
        period_end=date(2026, 2, 9),
        focus_name="Todos",
        attribution_df=attribution,
        selected_nps_df=payload["selected_nps"],
        comparison_nps_df=payload["comparison_nps"],
    )
    deck = Presentation(BytesIO(report.content))
    text = "\n".join(
        shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame
    )
    assert "CONFIANZA SEMÁNTICA" in text
    assert "NOTA MEDIA DE COMENTARIOS ENLAZADOS" in text
    assert "No demuestran causalidad" not in text
    for secret in ("12345678", "Acme", "Juan Pérez", "juan@example", "4444", "ABC999"):
        assert secret not in text
    service = object.__new__(DashboardService)
    cards = service._build_linking_scenario_cards(attribution)
    assert cards[0]["spotlight_metrics"][3]["label"] == link_confidence_label("llm")
    assert set(attribution.columns) <= set(cards[0])
    assert link_confidence_label("rules") == "SIMILITUD TEXTUAL"
    newsletter = build_executive_newsletter(
        current_df=payload["selected_nps"],
        period_kpis={},
        linking={"scenarios": {"cards": cards}},
        topic_channel="Web",
        period_start=date(2026, 1, 1),
        period_end=date(2026, 2, 9),
    )
    assert "Juan Pérez" not in json.dumps(newsletter, ensure_ascii=False)
    assert "Asociación operativa consistente" in newsletter["lead"]


def test_enrichment_preserves_literal_evidence_and_evaluates_operational_journey():
    links = pd.DataFrame(
        [
            link(
                nps_id=f"n{i}",
                incident_id=f"i{i % 2}",
                nps_topic="Acceso > Token",
                text_similarity=0.9,
                causal_engine="llm",
                affected_task="Generar token",
                observed_symptom="Pantalla en blanco",
            )
            for i in range(3)
        ]
    )
    nps = pd.DataFrame(
        {
            "ID": ["n0", "n1", "n2"],
            "Fecha": pd.to_datetime(["2026-09-03"] * 3),
            "NPS": [1, 2, 3],
            "Comment": ["No genera token"] * 3,
            "Palanca": ["Acceso"] * 3,
            "Subpalanca": ["Token"] * 3,
        }
    )
    helix = pd.DataFrame(
        {
            "Incident Number": ["i0", "i1"],
            "Submit Date": pd.to_datetime(["2026-09-01"] * 2),
            "Detailed Description": ["Error al generar token"] * 2,
        }
    )
    result = build_incident_attribution_chains(links, nps, helix, top_k=0)
    assert result.iloc[0]["evidence_level"] == "EVIDENCIA_OPERATIVA_RECURRENTE"
    assert result.iloc[0]["linked_comments"] == 3
    assert (
        "Generar token → Pantalla en blanco → 2 incidencias Helix"
        in result.iloc[0]["journey_route"]
    )


def test_reporting_rejects_selected_llm_engine_instead_of_silent_rules_fallback(
    exchange, monkeypatch
):
    handler, ctx, _, client = exchange
    service = client.app.state.dashboard_service
    state = handler.taxonomy.state(ctx)
    state["causal_engine"] = "llm"
    handler.taxonomy.save_state(ctx, state)
    monkeypatch.setattr(
        service,
        "classification_status",
        lambda *a, **kw: {
            "ready": False,
            "selected_engine": "llm",
            "reason": "Fingerprint desactualizado.",
        },
    )
    with pytest.raises(
        ReportCoherenceError, match="Regenera la clasificación de incidencias Helix"
    ):
        service.generate_ppt_report(
            context=ctx, pop_year="2026", pop_month="09", score_channel="Web"
        )


def test_journey_match_rejects_generic_and_tied_signals():
    from nps_lens.analytics.incident_attribution import _executive_journey_match

    args = dict(
        nps_topic="",
        touchpoint="",
        palanca="",
        subpalanca="",
        incident_topic="",
        incident_summary="",
        comment_txt="pago rechazado",
    )
    catalog = [{"name": "one", "keywords": ["pago"]}]
    assert _executive_journey_match(**args, catalog=catalog) is None
    catalog = [
        {"name": "one", "keywords": ["pago", "rechazado"]},
        {"name": "two", "keywords": ["pago", "rechazado"]},
    ]
    assert _executive_journey_match(**args, catalog=catalog) is None
    assert _executive_journey_match(**args, catalog=catalog[:1])["name"] == "one"


def test_links_assign_every_supported_topic_without_pair_inflation():
    from nps_lens.analytics.nps_helix_link import link_incidents_to_nps_topics, weekly_aggregates

    nps = pd.DataFrame(
        {
            "ID": ["a", "b", "c"],
            "_business_key": ["a", "b", "c"],
            "Comment": ["transferencias rechazadas"] * 3,
            "Palanca": ["Mayoritaria", "Mayoritaria", "Minoritaria"],
            "Subpalanca": ["Transferir"] * 3,
            "NPS": [2] * 3,
            "Fecha": pd.to_datetime(["2026-09-01"] * 3),
        }
    )
    helix = pd.DataFrame(
        {
            "Incident Number": ["i"],
            "summary": ["transferencias rechazadas"],
            "Fecha": pd.to_datetime(["2026-09-01"]),
        }
    )
    assignments, links = link_incidents_to_nps_topics(nps, helix)
    assert len(links) == 3 and len(assignments) == 2
    _, topics = weekly_aggregates(nps, helix, assignments)
    assert topics.incidents.tolist() == [1, 1]
