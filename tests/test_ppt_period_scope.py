from datetime import date
from io import BytesIO

import pandas as pd
import pytest
from pptx import Presentation

from nps_lens.reports import executive_ppt as ppt
from nps_lens.reports.coherence import ReportCoherenceError
from nps_lens.services.analytics.kpis_service import build_period_kpis


def history():
    return pd.DataFrame(
        {
            "Fecha": pd.to_datetime(
                ["2026-01-10", "2026-02-10", "2026-03-01", "2026-03-18 23:59", "2026-04-01"],
                format="mixed",
            ),
            "NPS": [0, 10, 0, 10, 0],
            "Comment": ["mal", "bien", "lento", "rápido", "fuera del periodo"],
            "Palanca": ["Servicio"] * 5,
            "Subpalanca": ["Atención"] * 5,
            "Canal": ["Web"] * 5,
        }
    )


def context(frame=None, **overrides):
    frame = history() if frame is None else frame
    args = dict(
        service_origin="Banco",
        service_origin_n1="",
        service_origin_n2="",
        period_start=date(2026, 3, 1),
        period_end=date(2026, 3, 18),
        focus_name="Todos",
        topic_channel="Web",
        selected_nps_df=frame,
        comparison_nps_df=frame,
        attribution_df=None,
        touchpoint_source="",
        entity_summary_df=None,
        entity_summary_kpis=None,
        broken_journeys_df=None,
    )
    args.update(overrides)
    return ppt._build_presentation_context(**args)


def test_historic_rates_use_cumulative_population_and_include_final_day():
    frame = history()
    selected = frame.loc[frame.Fecha.dt.month.eq(3)]
    kpis = build_period_kpis(
        history_df=frame,
        current_df=selected,
        pop_year="2026",
        pop_month="03",
        context_label="Marzo 2026",
    )
    result = context(period_kpis=kpis)
    assert result.overview["comments"] == 2
    assert result.overview["base_detractor_rate"] == 0.5
    assert result.overview["cumulative_detractor_rate"] == 0.5
    assert result.overview["start_detr"] == 0  # Monthly base is February, not historical base.
    assert result.overview["cumulative_classic_nps"] == 0
    assert result.causal.channel == "Web"
    assert all(
        pd.Timestamp(x) <= pd.Timestamp("2026-03-18") for x in result.overview_figure.data[0].x
    )


def test_internally_consistent_kpis_from_another_population_are_rejected():
    frame = history()
    wrong = build_period_kpis(
        history_df=frame,
        current_df=frame.iloc[[1]],
        pop_year="2026",
        pop_month="02",
        context_label="Febrero 2026",
    )
    with pytest.raises(ReportCoherenceError, match="Población del periodo"):
        context(period_kpis=wrong)


def test_company_scope_and_reversed_dates_are_rejected():
    frame = history().assign(service_origin="Otro banco")
    with pytest.raises(ReportCoherenceError, match="ámbito"):
        context(frame)
    with pytest.raises(ReportCoherenceError, match="inicio"):
        context(period_start=date(2026, 4, 1))


def test_voc_evidence_must_belong_to_period_but_incidents_can_precede_it():
    chain = pd.DataFrame([{"comment_records": [{"date": "28-02-2026", "comment": "fuera"}]}])
    with pytest.raises(ReportCoherenceError, match="evidencia VoC"):
        context(attribution_df=chain)
    result = context(evidence_channel="Todos")
    assert result.causal.channel == "Todos"


def test_multimonth_deck_has_no_monthly_or_fabricated_deterioration_claim(monkeypatch):
    monkeypatch.setattr(ppt, "_RENDERER_AVAILABLE", False)
    frame = history().iloc[:4]
    report = ppt.generate_business_review_ppt(
        service_origin="Banco",
        service_origin_n1="",
        service_origin_n2="",
        period_start=date(2026, 1, 1),
        period_end=date(2026, 3, 18),
        focus_name="Todos",
        topic_channel="Todos",
        selected_nps_df=frame,
        comparison_nps_df=history(),
        include_causal_section=False,
    )
    deck = Presentation(BytesIO(report.content))
    texts = "\n".join(
        shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame
    )
    assert "mensual" not in texts.lower()
    assert "del mes" not in texts.lower()
    assert "Sin deterioros comparables" in texts
    assert "Sin deterioro presenta" not in texts
    assert "fuera del periodo" not in texts
    assert "2026-01-01" in texts and "2026-03-18" in texts
    for slide in deck.slides:
        assert "Periodo VoC:" in slide.notes_slide.notes_text_frame.text


def test_selected_month_across_years_keeps_intervening_months_out_of_topic_metrics():
    frame = history().iloc[:3].copy()
    frame["Fecha"] = pd.to_datetime(["2025-03-01", "2025-04-01", "2026-03-01"])
    selected = frame.loc[frame.Fecha.dt.month.eq(3)]
    kpis = build_period_kpis(
        history_df=frame,
        current_df=selected,
        pop_year="Todos",
        pop_month="03",
        context_label="Marzo (todos los años)",
    )
    result = context(
        frame,
        selected_nps_df=selected,
        period_start=date(2025, 3, 1),
        period_end=date(2026, 3, 1),
        period_kpis=kpis,
    )
    assert result.overview["comments"] == 2
    assert result.dimensions["Palanca"].topic_table_df.iloc[0]["n"] == 2


def test_comparison_base_is_verified_against_history_not_just_its_own_formula():
    frame = history()
    current = frame.loc[frame.Fecha.dt.month.eq(3)]
    wrong = build_period_kpis(
        history_df=frame.loc[~frame.Fecha.dt.month.eq(2)],
        current_df=current,
        pop_year="2026",
        pop_month="03",
        context_label="Marzo 2026",
    )
    with pytest.raises(ReportCoherenceError, match="Población de la base"):
        context(frame, period_kpis=wrong)
