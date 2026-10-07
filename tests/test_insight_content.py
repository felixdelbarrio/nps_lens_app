from __future__ import annotations

import pandas as pd
import pytest

from nps_lens.analytics.drivers import driver_table
from nps_lens.analytics.signal_quality import actionable_rows, is_reserve_category, signal_quality
from nps_lens.domain.comment_text import useful_comment_mask
from nps_lens.reports import executive_ppt
from nps_lens.reports.executive_newsletter import _focus_rows
from nps_lens.services.analytics.kpis_service import compute_score_kpis
from nps_lens.services.dashboard_service import DashboardService


@pytest.mark.parametrize(
    "label",
    [
        "Sin comentarios",
        " SIN CONTENIDO. ",
        "Genérico",
        "GENÉRICA",
        "Comentario genérico",
        "Otros",
        "Valoración general del canal",
        "Experiencia global",
        "Sin clasificación temática > Información insuficiente",
        "Acceso / sin comentarios",
        "",
        None,
        pd.NA,
        float("nan"),
    ],
)
def test_placeholder_categories_never_qualify_as_insights(label):
    if label is not None and label is not pd.NA and pd.notna(label) and str(label):
        assert is_reserve_category(label)
    frame = pd.DataFrame({"Subpalanca": [label], "Comment": ["No puedo generar token"], "NPS": [0]})
    assert actionable_rows(frame).empty


@pytest.mark.parametrize(
    "label",
    [
        "Información sobre el estado de una transferencia",
        "Error genérico al firmar",
        "Sin acceso a la cuenta",
        "Faltan comentarios en el expediente",
    ],
)
def test_specific_topics_are_not_excluded_by_a_word_inside_their_name(label):
    assert not is_reserve_category(label)


def test_placeholder_comments_are_excluded_even_when_the_topic_name_looks_specific():
    frame = pd.DataFrame(
        {
            "Subpalanca": ["Token"] * 6,
            "Comment": ["sin comentarios", "...", pd.NA, "GENÉRICO", "", "No genera token"],
            "NPS": [0, 0, 10, 10, 10, 2],
        }
    )
    assert useful_comment_mask(frame).tolist() == [False] * 5 + [True]
    assert actionable_rows(frame).Comment.tolist() == ["No genera token"]
    result = executive_ppt._period_overview(frame)
    assert result["friction"]["n"] == 1
    assert result["strength"] == {}
    assert compute_score_kpis(frame).samples == 6


@pytest.mark.parametrize("categorical", [False, True])
def test_generic_topics_remain_in_nps_and_gaps_but_never_lead_editorial_rankings(categorical):
    frame = pd.DataFrame(
        {
            "Palanca": ["Genérico"] * 6 + ["Sin comentarios"] * 6 + ["Acceso", "Atención"],
            "Subpalanca": ["Genérico"] * 6 + ["Sin comentarios"] * 6 + ["Token", "Ayuda eficaz"],
            "Comment": ["opinión"] * 12 + ["No genera token", "Me ayudaron a resolver"],
            "NPS": [0] * 6 + [10] * 6 + [2, 10],
        }
    )
    if categorical:
        frame = frame.astype(
            {"Palanca": "category", "Subpalanca": "category", "Comment": "category"}
        )
    kpis = compute_score_kpis(frame)
    assert kpis.samples == 14 and kpis.classic_nps == 0
    gaps = driver_table(frame, "Subpalanca")
    assert {row.value for row in gaps} == {"Genérico", "Sin comentarios", "Token", "Ayuda eficaz"}
    assert sum(row.valid_n for row in gaps) == 14
    insights = executive_ppt._period_overview(frame)
    assert insights["friction"]["topic"] == "Token"
    assert insights["strength"]["topic"] == "Ayuda eficaz"
    assert signal_quality(frame)["insufficient_comments"] == 12
    assert {row["label"] for row in _focus_rows(frame, topic_channel="Todos")} == {
        "Acceso",
        "Atención",
    }
    service = DashboardService.__new__(DashboardService)
    clusters = service._topics_df(frame)
    assert {term for terms in clusters.top_terms for term in terms} == {"Token", "Ayuda eficaz"}


def test_no_meaningful_topic_produces_an_explicit_empty_insight():
    frame = pd.DataFrame(
        {"Subpalanca": ["Genérico", "Sin comentarios"], "Comment": ["opinión", ""], "NPS": [0, 10]}
    )
    result = executive_ppt._period_overview(frame)
    assert result["friction"] == {} and result["strength"] == {}
    for positive in (False, True):
        assert "No hay un tópico con comentarios útiles" in executive_ppt._topic_signal_copy(
            {}, positive=positive
        )


def test_newsletter_does_not_invent_a_generic_friction_when_only_empty_topics_remain():
    from datetime import date

    from nps_lens.reports.executive_newsletter import build_executive_newsletter
    from nps_lens.services.analytics.kpis_service import build_period_kpis

    frame = pd.DataFrame(
        {
            "Fecha": pd.to_datetime(["2026-08-01", "2026-08-02"]),
            "Palanca": ["Genérico", "Sin contenido"],
            "Comment": ["opinión", ""],
            "NPS": [0, 10],
        }
    )
    kpis = build_period_kpis(
        history_df=frame,
        current_df=frame,
        pop_year="2026",
        pop_month="08",
        context_label="Agosto 2026",
    )
    newsletter = build_executive_newsletter(
        current_df=frame,
        period_kpis=kpis,
        linking={},
        topic_channel="Todos",
        period_start=date(2026, 8, 1),
        period_end=date(2026, 8, 2),
    )
    assert newsletter["headline"] == "No hay evidencia suficiente para destacar un foco de fricción"
    assert newsletter["signals"] == []
    assert newsletter["scorecard"][1]["value"] == "0,00"


def test_ppt_does_not_turn_an_empty_topic_list_into_a_generic_conclusion(monkeypatch):
    from datetime import date
    from io import BytesIO

    from pptx import Presentation

    monkeypatch.setattr(executive_ppt, "_RENDERER_AVAILABLE", False)
    frame = pd.DataFrame(
        {
            "Fecha": pd.to_datetime(["2026-08-01", "2026-08-02"]),
            "Palanca": ["Genérico", "Sin contenido"],
            "Subpalanca": ["Sin comentarios"] * 2,
            "Comment": ["opinión", ""],
            "NPS": [0, 10],
        }
    )
    report = executive_ppt.generate_business_review_ppt(
        service_origin="Banco",
        service_origin_n1="",
        service_origin_n2="",
        period_start=date(2026, 8, 1),
        period_end=date(2026, 8, 2),
        focus_name="Todos",
        selected_nps_df=frame,
        comparison_nps_df=frame,
        include_causal_section=False,
    )
    deck = Presentation(BytesIO(report.content))
    comparison = deck.slides[2]
    assert "No hay un tópico con comentarios útiles" in " ".join(comparison.shapes[8].text.split())
    assert "No hay un tópico con comentarios útiles" in " ".join(comparison.shapes[12].text.split())
    topics = " ".join(shape.text for shape in deck.slides[3].shapes if shape.has_text_frame)
    assert "No hay comentarios detractores útiles para identificar temas" in topics
    assert "concentra su señal principal" not in topics
