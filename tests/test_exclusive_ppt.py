from __future__ import annotations

from datetime import date
from io import BytesIO

import pandas as pd
from pptx import Presentation

from nps_lens.reports.exclusive_ppt import (
    ExclusiveReportContext,
    find_exclusive_template_path,
    generate_exclusive_report,
)


def test_exclusive_report_deduplicates_opportunities_and_links_helix() -> None:
    rows = []
    labels = ["Pagos y transferencias", "Pagos/transferencias", "Login"]
    for month in [1, 2, 3]:
        for index in range(90):
            rows.append(
                {
                    "Fecha": pd.Timestamp(2026, month, (index % 20) + 1),
                    "NPS": index % 11,
                    "Palanca": labels[index % len(labels)],
                    "Subpalanca": "Fallas en el Login",
                    "Comment": "No pude completar el pago y necesité ayuda.",
                }
            )
    frame = pd.DataFrame(rows)
    attribution = pd.DataFrame(
        [
            {
                "touchpoint": "Pagos",
                "nps_topic": "Pagos/transferencias",
                "palanca": "Pagos/transferencias",
                "subpalanca": "Operativa",
                "linked_pairs": 8,
                "linked_incidents": 1,
                "linked_comments": 8,
                "confidence": 0.91,
                "incident_records": [
                    {
                        "incident_id": "INC000123",
                        "url": "https://helix.example/INC000123",
                    }
                ],
            }
        ]
    )
    result = generate_exclusive_report(
        template_path=find_exclusive_template_path(),
        context=ExclusiveReportContext(
            service_origin="BBVA México",
            service_origin_n1="ENTERPRISE WEB",
            service_origin_n2="",
            period_start=date(2026, 1, 1),
            period_end=date(2026, 3, 31),
            helix_base_url="https://helix.example/",
        ),
        selected_nps_df=frame,
        comparison_nps_df=frame,
        attribution_df=attribution,
        min_n=20,
    )
    presentation = Presentation(BytesIO(result.content))
    opportunity_text = "\n".join(
        shape.text for shape in presentation.slides[6].shapes if shape.has_text_frame
    )
    assert opportunity_text.casefold().count("pagos/transferencias") <= 1
    helix_runs = [
        run
        for shape in presentation.slides[9].shapes
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
        if run.text == "INC000123"
    ]
    assert helix_runs
    assert helix_runs[0].hyperlink.address == "https://helix.example/INC000123"
